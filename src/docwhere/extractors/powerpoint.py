"""PowerPoint: one chunk per slide (text boxes, tables, groups and speaker notes)."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

from docwhere.extractors import register
from docwhere.model import Chunk, norm


@register(".pptx", ".pptm")
def extract(path: Path) -> list[Chunk]:
    from pptx import Presentation

    chunks: list[Chunk] = []
    for no, slide in enumerate(Presentation(str(path)).slides, 1):
        title_shape = slide.shapes.title
        title = norm(title_shape.text_frame.text) if title_shape is not None else ""
        texts = list(_shape_texts(slide.shapes))
        if slide.has_notes_slide:
            texts += slide.notes_slide.notes_text_frame.text.splitlines()
        lines = [line for line in map(norm, texts) if line]
        if lines:
            chunks.append(Chunk(title, title, lines, [{"slide": no}]))
    return chunks


def _shape_texts(shapes) -> Iterator[str]:
    from pptx.shapes.group import GroupShape

    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _shape_texts(shape.shapes)
        elif shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                yield para.text
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text
