"""Excel: one line per row, so a hit carries the rest of its row and the column headers.

Cells in a line are tab-separated (``norm`` removes tabs from cell text, so the separator
never collides), and each line's anchor lists the column letter of every cell in it. The
sheet's header row, when one is found, opens every chunk of that sheet so a hit can be shown
as "コード: E2041 / 対処: ...".
"""

from __future__ import annotations

from pathlib import Path

from docwhere.extractors import register
from docwhere.model import Chunk, norm

ROWS_PER_CHUNK = 100
HEADER_SEARCH_ROWS = 10  # a header row is looked for among the first non-empty rows only
SEP = "\t"


@register(".xlsx", ".xlsm")
def extract(path: Path) -> list[Chunk]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[Chunk] = []
    try:
        for ws in wb.worksheets:
            chunks += _sheet_chunks(ws)
    finally:
        wb.close()
    return chunks


def _sheet_chunks(ws) -> list[Chunk]:
    chunks: list[Chunk] = []
    header: tuple[str, dict] | None = None
    looking_for_header = True
    seen_rows = 0
    lines: list[str] = []
    anchors: list[dict] = []

    def emit() -> None:
        if any(not a.get("header") for a in anchors):
            chunks.append(Chunk(ws.title, ws.title, lines[:], anchors[:]))

    for row in ws.iter_rows():
        filled = [(c, text) for c in row if c.value is not None and (text := norm(str(c.value)))]
        if not filled:
            continue
        cells = [(c.column_letter, c.value, text) for c, text in filled]
        seen_rows += 1
        line = SEP.join(text for _, _, text in cells)
        anchor = {"sheet": ws.title, "row": filled[0][0].row, "cols": [col for col, _, _ in cells]}

        if looking_for_header:
            if len(cells) >= 2 and all(isinstance(value, str) for _, value, _ in cells):
                header = (line, {**anchor, "header": True})
                looking_for_header = False
                lines.append(header[0])
                anchors.append(header[1])
                continue
            # a title row (one cell) may precede the header; numbers mean data has started
            if len(cells) >= 2 or seen_rows >= HEADER_SEARCH_ROWS:
                looking_for_header = False

        lines.append(line)
        anchors.append(anchor)
        if len(lines) >= ROWS_PER_CHUNK:
            emit()
            lines, anchors = ([header[0]], [header[1]]) if header else ([], [])
    emit()
    return chunks
