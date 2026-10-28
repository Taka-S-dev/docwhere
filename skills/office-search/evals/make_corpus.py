"""Build a realistic document folder for evaluating the docwhere skill.

The fictional OrbitSync CLI is documented across a current spec, an outdated spec, an ops
manual that only references the spec, a trouble-shooting deck, an error-code sheet, a generic
network guide, meeting notes, unrelated admin files and a scanned (text-less) PDF.
"""

import sys
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation


def word(path: Path, sections: list[tuple[int, str, list]]) -> None:
    doc = Document()
    for level, heading, body in sections:
        if level == 0:
            doc.add_heading(heading, 0)
        else:
            doc.add_heading(heading, level)
        for item in body:
            if isinstance(item, list):
                table = doc.add_table(rows=len(item), cols=len(item[0]))
                for r, row in enumerate(item):
                    for c, value in enumerate(row):
                        table.cell(r, c).text = value
            else:
                doc.add_paragraph(item)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def pdf(path: Path, pages: list[str], toc: list[list]) -> None:
    doc = pymupdf.open()
    for text in pages:
        doc.new_page().insert_htmlbox(pymupdf.Rect(50, 50, 545, 790), text.replace("\n", "<br>"))
    doc.set_toc(toc)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main(root: Path) -> None:
    spec = root / "仕様書" / "OrbitSync_CLI仕様書_v2.3.docx"
    word(spec, [
        (0, "OrbitSync CLI 仕様書 v2.3", ["発行: 基盤チーム 2026-08-01"]),
        (1, "1 概要", ["OrbitSync は社内ファイルサーバーとクラウドストレージを同期するコマンドラインツールである。",
                       "本書は orbitsync コマンドの書式、オプション、既定値を定める。"]),
        (1, "2 インストールと初期設定", []),
        (2, "2.1 動作環境", ["Windows 10 以降、macOS 13 以降。.NET ランタイム不要の単一実行ファイルで配布する。"]),
        (2, "2.2 ネットワーク設定", [
            "社内ネットワークからクラウドへ接続する場合は、プロキシサーバーを経由する必要がある。",
            "環境変数 HTTPS_PROXY にプロキシの URL を設定すること。例: HTTPS_PROXY=http://proxy.corp.example:8080",
            "認証付きプロキシの場合は orbitsync config set proxy.auth <ユーザー名> で資格情報を登録する。",
            "通信タイムアウトは設定キー net.timeout で変更できる（既定 30 秒、最大 300 秒）。",
        ]),
        (1, "3 コマンドリファレンス", ["各コマンドは orbitsync <コマンド> [オプション] の形式で実行する。"]),
        (2, "3.1 sync コマンド", [
            "書式: orbitsync sync [--dry-run] [--exclude <パターン>] [--target <リモート名>]",
            "ローカルとリモートの差分を双方向に同期する。",
            [["オプション", "説明", "既定値"], ["--dry-run", "変更内容を表示するだけで実行しない", "無効"],
             ["--exclude", "同期対象から除外するパス（glob）", "なし"], ["--target", "同期先のリモート名", "default"]],
        ]),
        (2, "3.2 config コマンド", [
            "書式: orbitsync config get|set|list <キー> [<値>]",
            "設定ファイル orbitsync.toml の値を参照・変更する。主なキー: net.timeout, proxy.auth, purge.retention_days",
        ]),
        (2, "3.3 purge コマンド", [
            "書式: orbitsync purge [--older-than <日数>] [--dry-run] [--target <リモート名>]",
            "リモートに保存された古い世代（スナップショット）を削除する。v2.0 で clean コマンドから改名した。",
            [["オプション", "説明", "既定値"],
             ["--older-than", "指定日数より古い世代を削除する", "90 日（purge.retention_days で変更可）"],
             ["--dry-run", "削除対象を一覧表示するだけで削除しない", "無効"],
             ["--target", "対象のリモート名", "default"]],
            "削除した世代は復元できない。初回は必ず --dry-run で対象を確認すること。",
            "削除対象が 1 件もない場合はエラー E4100 を返す。",
        ]),
        (2, "3.4 status コマンド", ["書式: orbitsync status", "最後の同期日時、未同期ファイル数、リモートの使用容量を表示する。"]),
        (1, "4 エラーコード", ["エラーコードの一覧と対処方法は別紙「エラーコード一覧.xlsx」を参照すること。"]),
        (1, "付録 A 変更履歴", [
            "v2.3 (2026-08): purge の既定保存期間を 30 日から 90 日に変更。",
            "v2.0 (2025-04): clean コマンドを purge に改名。",
        ]),
    ])

    pdf(root / "旧資料" / "OrbitSync_CLI仕様書_v1.4.pdf", [
        "OrbitSync CLI 仕様書 v1.4\n1 概要\nOrbitSync は社内ファイル同期ツールである。",
        "3 コマンドリファレンス\n3.1 sync コマンド\n書式: orbitsync sync [--dry-run]",
        "3.3 clean コマンド\n書式: orbitsync clean [--days <日数>]\n古い世代を削除する。--days の既定値は 30 日。\n"
        "注: 次期バージョンで purge に改名予定。",
    ], [[1, "1 概要", 1], [1, "3 コマンドリファレンス", 2], [2, "3.1 sync コマンド", 2], [2, "3.3 clean コマンド", 3]])

    word(root / "運用" / "月次運用手順書.docx", [
        (0, "OrbitSync 月次運用手順書", []),
        (1, "1 月次作業一覧", ["(1) 同期状況の確認 (2) 世代整理 (3) バックアップ確認"]),
        (1, "2 世代整理（purge）", [
            "毎月最終営業日に orbitsync purge --older-than 90 --dry-run で対象を確認し、問題なければ --dry-run を外して実行する。",
            "purge のオプションの詳細は CLI 仕様書 3.3 を参照。",
        ]),
        (1, "3 バックアップ確認", ["orbitsync status で最終同期日時が当日であることを確認する。"]),
    ])

    prs = Presentation()
    slides = [
        ("障害対応の流れ", "1. エラーコードを確認 2. エラーコード一覧で対処を確認 3. 解決しなければ基盤チームへ", ""),
        ("認証エラー（E2041）", "原因: 認証トークンの有効期限切れ（90 日）\n対処: orbitsync login --refresh を実行して再認証する",
         "SSO のパスワード変更直後にも発生する。再認証しても直らない場合はトークンキャッシュ削除。"),
        ("タイムアウト（E3002）", "社内ネットワークではプロキシ設定漏れが最も多い原因\n仕様書 2.2 の HTTPS_PROXY 設定を確認",
         ""),
        ("連絡先", "基盤チーム 内線 4321 / #orbitsync-support", ""),
    ]
    for title, body, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    (root / "運用").mkdir(parents=True, exist_ok=True)
    prs.save(root / "運用" / "障害対応手順_2026.pptx")

    wb = Workbook()
    ws = wb.active
    ws.title = "エラーコード"
    ws.append(["コード", "分類", "内容", "対処"])
    for row in [
        ("E1001", "設定", "設定ファイル orbitsync.toml の書式不正", "orbitsync config list で該当キーを確認"),
        ("E2041", "認証", "認証トークンの有効期限切れ", "orbitsync login --refresh で再認証"),
        ("E2050", "認証", "権限不足", "管理者にリモートの書き込み権限を依頼"),
        ("E3002", "通信", "接続タイムアウト", "プロキシ設定（HTTPS_PROXY）を確認。回線が遅い場合は net.timeout を延長"),
        ("E4100", "purge", "削除対象の世代がない", "対処不要（--older-than の日数を確認）"),
    ]:
        ws.append(row)
    wb.create_sheet("改訂履歴").append(["2026-08-01", "E4100 を追加"])
    (root / "管理").mkdir(parents=True, exist_ok=True)
    wb.save(root / "管理" / "エラーコード一覧.xlsx")

    pdf(root / "ネットワーク" / "社内ネットワーク接続ガイド.pdf", [
        "社内ネットワーク接続ガイド\n1 社内ネットワーク概要\n社内から外部への通信はすべてプロキシサーバーを経由する。直接接続はファイアウォールで遮断される。",
        "2 プロキシ設定\n2.1 ブラウザ\n自動構成スクリプト http://proxy.corp.example/proxy.pac を指定する。",
        "2.2 コマンドラインツール\n多くの CLI ツールは環境変数でプロキシを指定する。\n"
        "HTTPS_PROXY=http://proxy.corp.example:8080\nNO_PROXY=localhost,.corp.example\n設定しないと外部接続がタイムアウトする。",
        "3 VPN\n在宅勤務時は VPN 接続後に上記プロキシ設定がそのまま有効になる。",
    ], [[1, "1 社内ネットワーク概要", 1], [1, "2 プロキシ設定", 2], [2, "2.1 ブラウザ", 2],
        [2, "2.2 コマンドラインツール", 3], [1, "3 VPN", 4]])

    (root / "議事録").mkdir(parents=True, exist_ok=True)
    (root / "議事録" / "2026-09-15_基盤定例.md").write_text(
        "# 基盤チーム定例 2026-09-15\n\n## 決定事項\n- purge の既定保存期間 90 日化は v2.3 で対応済み。\n"
        "- E3002（タイムアウト）の問い合わせが多いため、社内ネットワーク接続ガイドを再周知する。\n\n"
        "## 次回\n- 10 月定例で status コマンドの出力改善を議論。\n",
        encoding="utf-8",
    )

    word(root / "総務" / "経費精算マニュアル.docx", [
        (0, "経費精算マニュアル", []),
        (1, "1 申請期限", ["経費は発生月の翌月 5 営業日までに申請する。期限を過ぎた申請はエラーとして差し戻される。"]),
        (1, "2 交通費", ["定期区間内の交通費は精算できない。"]),
    ])
    wb = Workbook()
    wb.active.title = "3F"
    wb.active.append(["席番号", "氏名", "内線"])
    wb.active.append(["3F-12", "基盤チーム", "4321"])
    wb.save(root / "総務" / "座席表.xlsx")

    scan = pymupdf.open()
    page = scan.new_page()
    page.draw_rect(pymupdf.Rect(100, 100, 400, 300), color=(0, 0, 0), fill=(0.9, 0.9, 0.9))
    scan.save(root / "旧資料" / "scan_設計メモ_2024.pdf")


def legacy_and_layout(root: Path) -> None:
    """A legacy .doc and a PDF without bookmarks, each the only source of one answer."""
    import win32com.client

    tmp = root / "運用" / "_tmp_log.docx"
    word(tmp, [
        (0, "OrbitSync ログ運用手順", ["作成: 運用チーム 2025-02-10"]),
        (1, "1 目的", ["障害調査に必要なログを確実に残すための手順を定める。"]),
        (1, "2 ログの種類", ["同期ログ、監査ログ、デバッグログの 3 種類がある。"]),
        (1, "3 ログの保存先と保存期間", [
            "ログは C:\\ProgramData\\OrbitSync\\logs に日付ごとのファイルで出力される。",
            "保存期間は 14 日。それより古いファイルは毎日 3:00 に自動削除される。",
            "保存先は設定キー log.dir で変更できる。",
        ]),
        (1, "4 障害時の採取", ["orbitsync diag --collect で直近 3 日分のログを zip にまとめる。"]),
    ])
    app = win32com.client.DispatchEx("Word.Application")
    try:
        doc = app.Documents.Open(str(tmp.resolve()))
        doc.SaveAs2(str((root / "運用" / "ログ運用手順.doc").resolve()), FileFormat=0)
        doc.Close(0)
    finally:
        app.Quit()
    tmp.unlink()

    doc = pymupdf.open()
    pages = [
        "<h1>OrbitSync 性能チューニングガイド</h1><p>本書は大容量の同期を行う環境向けの調整方法をまとめる。</p>"
        "<h2>1 同期の並列数</h2><p>設定キー sync.workers で同時に転送するファイル数を指定する。既定は 4。</p>",
        "<h2>2 帯域制限</h2><p>業務時間帯に回線を圧迫しないよう、設定キー net.bandwidth_limit で上り帯域の上限を"
        "指定できる。単位は Mbps、既定は 0（無制限）。</p><p>時間帯ごとに変える場合は net.bandwidth_schedule に"
        "「09:00-18:00=20」の形式で記述する。</p>"
        "<h2>3 大きなファイル</h2><p>1 GB を超えるファイルは 64 MB ごとに分割して転送する。</p>",
    ]
    for html in pages:
        doc.new_page().insert_htmlbox(pymupdf.Rect(50, 50, 545, 790), html)
    (root / "技術資料").mkdir(parents=True, exist_ok=True)
    doc.save(root / "技術資料" / "OrbitSync性能チューニングガイド.pdf")


TOPICS = ["経費精算", "勤怠管理", "会議室予約", "VPN 申請", "PC 貸与", "名刺発注", "健康診断", "研修計画",
          "入館証", "郵便物", "備品購入", "出張申請", "在宅勤務", "防災訓練", "社内報", "採用面接"]


def noise(root: Path, count: int) -> None:
    """Unrelated and loosely related documents, some mentioning OrbitSync in passing."""
    import random

    rng = random.Random(42)
    for i in range(count):
        topic = rng.choice(TOPICS)
        dept = rng.choice(["総務", "人事", "経理", "情報システム", "営業", "企画"])
        mention = "共有フォルダは OrbitSync で同期されている。設定やエラーの詳細は基盤チームの資料を参照。" if i % 6 == 0 else ""
        paras = [f"{topic}に関する{dept}の手順。申請は社内ポータルから行う。", f"問い合わせは{dept}窓口まで。", mention]
        kind = i % 4
        folder = root / "部署資料" / dept
        folder.mkdir(parents=True, exist_ok=True)
        name = f"{topic}_{i:03d}"
        if kind == 0:
            word(folder / f"{name}.docx", [(0, f"{topic}手順", []), (1, "1 概要", paras), (1, "2 注意事項", ["期限を守ること。"])])
        elif kind == 1:
            wb = Workbook()
            wb.active.append(["項目", "内容", "担当"])
            for j in range(20):
                wb.active.append([f"{topic}{j}", rng.choice(paras[:2]), dept])
            if mention:
                wb.active.append(["同期", mention, "基盤"])
            wb.save(folder / f"{name}.xlsx")
        elif kind == 2:
            p = pymupdf.open()
            p.new_page().insert_htmlbox(pymupdf.Rect(50, 50, 545, 790), f"<h1>{topic}</h1>" + "".join(f"<p>{x}</p>" for x in paras if x))
            p.save(folder / f"{name}.pdf")
        else:
            (folder / f"{name}.md").write_text(f"# {topic}\n\n" + "\n\n".join(x for x in paras if x) + "\n", encoding="utf-8")


if __name__ == "__main__":
    target = Path(sys.argv[1])
    main(target)
    if len(sys.argv) > 2:
        legacy_and_layout(target)
        noise(target, int(sys.argv[2]))
